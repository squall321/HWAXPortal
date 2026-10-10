# Word 정리본이 LLM 에 넘기는 대화 원문을 모델 창에서 유도하는지 — 60,000자 고정에서 심의 대화의 결정문이 잘렸다
"""정리본(report) 모드는 대화 원문을 에이전트 서버 /chat 에 실어 LLM 이 보고서로 재구성하게 한다. 원문을 머리 60,000자에서
잘랐는데(docx_export.transcript_text), 메시지마다 좌석 발언을 먼저·결정문을 뒤에 쌓으므로 심의 대화에서는 잘리는 쪽이
**결정문**이었다. 모델은 '(…이후 생략)' 을 봤지만 받는 사람의 문서에는 아무 말이 없었다. 60,000 은 16K·128K 창 시절 값이다 —
운영 창은 1,048,576 토큰이다(docs/risk-review-redesign D-21).

이제 상한은 에이전트 서버가 알려 주는 모델 창(/health 의 context_tokens)에서 유도하고, 창을 못 읽으면 종전 60,000자(바닥)로
돈다. 잘랐으면 문서 머리와 감사 기록에 원문 길이를 남긴다. 좌석 발언은 화면 버블(say)이 아니라 줄이지 않은 글(sayFull)을 쓴다.
"""
import io
import json

import httpx
import pytest

from app.agent import docx_export as dx


def _conv(turn_say: str, turn_full: str | None, decision: str) -> dict:
    t = {"round": 1, "persona": "mech-a", "say": turn_say}
    if turn_full is not None:
        t["sayFull"] = turn_full
    return {"title": "심의", "messages": [{"role": "user", "text": "/심의 x"},
                                         {"role": "assistant", "text": decision, "delib": {"turns": [t]}}]}


# ── 특성화 — 고치기 전 코드에서도 초록이어야 한다 ─────────────────────────────────────────────
def test_짧은_대화는_종전대로_전부_넘긴다():
    tx = dx.transcript_text(_conv("버블", None, "결정문"))
    assert tx == "[질문] /심의 x\n\n[mech-a R1] 버블\n\n[응답] 결정문"


# ── 재현 — 고치기 전 코드에서 붉다 ───────────────────────────────────────────────────────────
def test_좌석_발언은_줄이지_않은_글을_쓴다():
    tx = dx.transcript_text(_conv("버블", "버블이 아니라 줄이지 않은 발언 전체", "결정문"))
    assert "[mech-a R1] 버블이 아니라 줄이지 않은 발언 전체" in tx


def test_상한을_비우면_자르지_않는다():
    decision = "(1) 결론 " + "가" * 90_000 + " (7) 신뢰도 — 중간"
    tx = dx.transcript_text(_conv("버블", None, decision), limit=None)
    assert tx.endswith("(7) 신뢰도 — 중간"), "결정문 뒤쪽이 잘렸다"


def test_인자_없이_부르는_옛_호출부는_종전_상한을_받는다():
    """배포 중 보류된 옛 포털 프로세스는 이 모듈을 함수 안에서 처음 쓸 때 import 하고 `transcript_text(conv)` 로 부른다(origin/main
    routes.py). 바인드 마운트라 그때 집는 것은 새 파일이다 — 기본값이 '무제한' 이면 그 섞인 판만 상한도 표식도 없이 원문 전체를
    에이전트로 보낸다(검토 P-6). 기본값은 종전 60,000자로 두고 새 호출부가 상한을 명시한다."""
    decision = "가" * 70_000
    tx = dx.transcript_text(_conv("버블", None, decision))
    assert "중 앞 60,000자만 전달했다" in tx and len(tx) < 61_000
    import inspect

    from app.agent import routes as R

    src = inspect.getsource(R._export_docx_inner)
    assert "dx.transcript_text, conv, None" in src, "새 호출부는 상한을 명시한다(창에서 유도한 값으로 따로 자른다)"


def test_잘랐으면_원문_길이를_표식에_넣는다():
    conv = _conv("버블", None, "가" * 1000)
    full = dx.transcript_text(conv)
    tx = dx.transcript_text(conv, limit=100)
    assert tx.startswith(full[:100]) and f"원문 {len(full):,}자 중 앞 100자" in tx


def test_전문_내보내기도_줄이지_않은_발언을_쓴다():
    from docx import Document

    data = dx.build_transcript(_conv("버블", "줄이지 않은 발언 전체 — 뒤쪽 인용 [e:3]", "결정문"))
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "뒤쪽 인용 [e:3]" in text


def test_요청_모델이_sayFull_을_받고_총량에_센다():
    from pydantic import ValidationError

    from app.agent.routes import DocxRequest

    body = DocxRequest(conversation=_conv("버블", "전체", "결정문"))
    assert body.conversation.messages[1].delib.turns[0].sayFull == "전체"
    with pytest.raises(ValidationError, match="총량"):
        DocxRequest(conversation=_conv("버블", "가" * 3_999_990, "결정문 열두 자보다 길다"))


# ── 라우트 — 창에서 유도한 상한과 사람에게 남는 표식 ──────────────────────────────────────────
@pytest.fixture()
def export():
    """정리본 내보내기 → (에이전트 서버가 받은 /chat 본문, 감사 meta, 문서 글)."""
    from docx import Document
    from fastapi.testclient import TestClient

    from app.auth.provider import Principal
    from app.deps import principal_pat_or_session
    from app.main import app

    app.dependency_overrides[principal_pat_or_session] = lambda: Principal(
        subject="u-docx", email="u@hwax.local", display_name="U", groups=["feat:deliberation"])

    def go(conv: dict, health: dict | None) -> tuple[str, dict, str]:
        sent: list[str] = []

        def upstream(req: httpx.Request) -> httpx.Response:
            if req.url.path.endswith("/health"):
                return httpx.Response(200, json=health) if health is not None else httpx.Response(503)
            sent.append(json.loads(req.content)["message"])
            body = 'event: result\ndata: {"content": "# 정리본 제목\\n## 요약\\n본문"}\n\nevent: done\ndata: {}\n\n'
            return httpx.Response(200, content=body.encode(), headers={"content-type": "text/event-stream"})

        with TestClient(app) as c:
            real = app.state.agent_client
            app.state.agent_client = httpx.AsyncClient(transport=httpx.MockTransport(upstream))
            try:
                r = c.post("/agent/export/docx", json={"conversation": conv, "mode": "report"})
            finally:
                app.state.agent_client = real
            row = app.state.agent_audit._conn.execute(
                "SELECT meta FROM agent_audit WHERE principal = 'u-docx' AND event = 'docx_report' ORDER BY id DESC LIMIT 1"
            ).fetchone()
        assert r.status_code == 200, r.text
        assert len(sent) == 1
        text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
        return sent[0], json.loads(row[0]), text

    try:
        yield go
    finally:
        app.dependency_overrides.pop(principal_pat_or_session, None)


def test_1M_창이면_결정문을_자르지_않고_넘긴다(export):
    """**이 시험이 이 파일의 이유다** — 75,000자 결정문(실측 최대급)이 60,000자에서 잘렸다."""
    decision = "(1) 결론 " + "브래킷 두께 0.6 mm 에서 힌지 응력이 항복을 넘는다. " * 2500 + "(10) 한계 — 끝"
    msg, meta, text = export(_conv("버블", None, decision), {"context_tokens": 1_048_576})
    assert msg.endswith("(10) 한계 — 끝"), "결정문 뒤쪽이 정리본에 가지 않았다"
    assert "생략" not in msg
    assert meta["tx_chars"] == meta["tx_sent"] and meta["window"] == "server"
    assert "앞부분만" not in text


def test_창을_못_읽으면_종전_상한으로_돌고_문서와_감사에_남긴다(export):
    decision = "가" * 70_000
    msg, meta, text = export(_conv("버블", None, decision), None)
    assert "원문" in msg and "중 앞 60,000자" in msg, "바닥은 종전 값이다"
    assert meta["window"] == "fallback" and meta["tx_sent"] == 60_000 and meta["tx_chars"] > 70_000
    assert f"원문 {meta['tx_chars']:,}자 중 앞 60,000자" in text, "받는 사람의 문서에 정리에 쓴 범위가 없다"


def test_작은_창에서도_종전_바닥_아래로_줄이지_않는다(export):
    """dev 16K 창 — 창에서 유도한 값이 종전 60,000자보다 작으면 종전 값을 쓴다(종전 동작이 깨지지 않게)."""
    _msg, meta, _text = export(_conv("버블", None, "가" * 70_000), {"context_tokens": 16_384})
    assert meta["tx_sent"] == 60_000 and meta["window"] == "server"


def test_총량은_렌더하는_칸만_센다():
    """렌더(_say)는 sayFull 이 있으면 say(버블)를 쓰지 않는다. 둘 다 세면 버블이 이중으로 세어져, 렌더할 글이 400만 자 안인데도
    422 다(검토 P-4). 좌석 2,000건 × (발언 1,900자 + 버블 600자) — 렌더는 380만 자, 둘 다 세면 500만 자."""
    from app.agent.routes import DocxRequest

    turn = {"round": 1, "persona": "mech-a", "say": "버" * 600, "sayFull": "발" * 1_900}
    msgs = [{"role": "user", "text": "/심의 x"}] + [
        {"role": "assistant", "text": "결정문", "delib": {"turns": [dict(turn) for _ in range(400)]}} for _ in range(5)]
    body = DocxRequest.model_validate({"conversation": {"title": "심의", "messages": msgs}, "mode": "transcript"})
    t0 = body.conversation.messages[1].delib.turns[0]
    assert t0.sayFull == "발" * 1_900 and dx._say(t0.model_dump()) == "발" * 1_900, "렌더할 글은 그대로다"


def test_버블만_있는_턴은_종전대로_say_를_센다():
    from pydantic import ValidationError

    from app.agent.routes import DocxRequest

    with pytest.raises(ValidationError, match="총량"):
        DocxRequest(conversation=_conv("가" * 3_999_990, None, "결정문 열두 자보다 길다"))
